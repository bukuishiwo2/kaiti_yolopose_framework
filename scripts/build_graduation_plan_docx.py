#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
GENERATOR_VERSION = "2.0.0"
DEFAULT_SOURCE = PROJECT_ROOT / "docs/graduation_research_plan_v2.md"
DEFAULT_OUTPUT = PROJECT_ROOT / "docs/generated/毕业设计研究内容与实施方案_叶豪成_修订版v2.0.docx"
DEFAULT_MANIFEST = PROJECT_ROOT / "docs/generated/graduation_research_plan_v2.manifest.json"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the graduation research plan DOCX and manifest.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--version", action="store_true", help="Print generator version and exit.")
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def relpath(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path.resolve())


def git_metadata() -> tuple[str | None, bool | None]:
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=PROJECT_ROOT,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()
        dirty = bool(
            subprocess.run(
                ["git", "status", "--porcelain"],
                cwd=PROJECT_ROOT,
                check=True,
                capture_output=True,
                text=True,
            ).stdout.strip()
        )
        return commit, dirty
    except (OSError, subprocess.CalledProcessError):
        return None, None


def set_run_font(run, *, east_asia: str, latin: str, size: float, bold: bool | None = None) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return text.strip()


def split_table_row(line: str) -> list[str]:
    return [clean_inline(cell) for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def iter_blocks(lines: list[str]) -> Iterable[tuple[str, object]]:
    index = 0
    paragraph: list[str] = []

    def flush_paragraph() -> tuple[str, object] | None:
        if not paragraph:
            return None
        value = " ".join(part.strip() for part in paragraph if part.strip())
        paragraph.clear()
        return ("paragraph", value)

    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("```"):
            block = flush_paragraph()
            if block:
                yield block
            language = line[3:].strip()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].rstrip().startswith("```"):
                code_lines.append(lines[index].rstrip("\n"))
                index += 1
            yield "code", {"language": language, "text": "\n".join(code_lines)}
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            block = flush_paragraph()
            if block:
                yield block
            yield "heading", {"level": len(heading.group(1)), "text": clean_inline(heading.group(2))}
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            block = flush_paragraph()
            if block:
                yield block
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            yield "table", rows
            continue

        if re.match(r"^\s*-\s+", line):
            block = flush_paragraph()
            if block:
                yield block
            yield "bullet", clean_inline(re.sub(r"^\s*-\s+", "", line))
            index += 1
            continue

        numbered = re.match(r"^\s*(\d+)[.)]\s+(.+)$", line)
        if numbered:
            block = flush_paragraph()
            if block:
                yield block
            yield "number", {
                "ordinal": int(numbered.group(1)),
                "text": clean_inline(numbered.group(2)),
            }
            index += 1
            continue

        if line.startswith(">"):
            block = flush_paragraph()
            if block:
                yield block
            yield "quote", clean_inline(line.lstrip("> "))
            index += 1
            continue

        if not line.strip():
            block = flush_paragraph()
            if block:
                yield block
        else:
            paragraph.append(line)
        index += 1

    block = flush_paragraph()
    if block:
        yield block


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, east_asia, size in [
        ("Title", "黑体", 18),
        ("Heading 1", "黑体", 16),
        ("Heading 2", "黑体", 14),
        ("Heading 3", "黑体", 12),
    ]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
        style.font.size = Pt(size)
        style.font.bold = True


def add_text_paragraph(document: Document, text: str, *, style: str | None = None, indent: bool = True):
    paragraph = document.add_paragraph(style=style)
    if not indent:
        paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(clean_inline(text))
    set_run_font(run, east_asia="宋体", latin="Times New Roman", size=12)
    return paragraph


def build_document(source: Path, output: Path) -> tuple[str, int]:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    title = "毕业设计研究内容与实施方案"
    block_count = 0

    for kind, payload in iter_blocks(lines):
        block_count += 1
        if kind == "heading":
            level = int(payload["level"])
            text = str(payload["text"])
            if level == 1:
                title = text
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                run = paragraph.add_run(text)
                set_run_font(run, east_asia="黑体", latin="Times New Roman", size=18, bold=True)
            else:
                paragraph = document.add_heading(text, level=min(level - 1, 3))
                paragraph.paragraph_format.first_line_indent = Pt(0)
        elif kind == "paragraph":
            add_text_paragraph(document, str(payload))
        elif kind == "bullet":
            add_text_paragraph(document, str(payload), style="List Bullet", indent=False)
        elif kind == "number":
            add_text_paragraph(
                document,
                f"{int(payload['ordinal'])}. {payload['text']}",
                indent=False,
            )
        elif kind == "quote":
            paragraph = add_text_paragraph(document, str(payload), indent=False)
            paragraph.paragraph_format.left_indent = Cm(0.8)
        elif kind == "code":
            paragraph = document.add_paragraph(style="No Spacing")
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(str(payload["text"]))
            set_run_font(run, east_asia="等线", latin="Courier New", size=9)
        elif kind == "table":
            rows = list(payload)
            if not rows:
                continue
            column_count = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=column_count)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index in range(column_count):
                    value = row[col_index] if col_index < len(row) else ""
                    cell = table.cell(row_index, col_index)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    run = paragraph.add_run(value)
                    set_run_font(
                        run,
                        east_asia="黑体" if row_index == 0 else "宋体",
                        latin="Times New Roman",
                        size=9.5,
                        bold=row_index == 0,
                    )
                    if row_index == 0:
                        set_cell_shading(cell, "D9EAF7")

    footer = document.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    footer_run = footer.add_run("研究方案 v2.0 · 由 Markdown 自动生成")
    set_run_font(footer_run, east_asia="宋体", latin="Times New Roman", size=9)

    document.core_properties.title = title
    document.core_properties.subject = "毕业设计研究内容与实施方案"
    document.core_properties.author = "叶豪成"
    document.core_properties.comments = (
        "Generated from docs/graduation_research_plan_v2.md; do not edit the DOCX manually."
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return title, block_count


def main() -> None:
    args = parse_args()
    if args.version:
        print(GENERATOR_VERSION)
        return

    source = args.source.resolve()
    output = args.output.resolve()
    manifest_path = args.manifest.resolve()
    if not source.exists():
        raise SystemExit(f"Research plan source does not exist: {source}")

    source_hash = sha256(source)
    commit, dirty = git_metadata()
    title, block_count = build_document(source, output)
    output_hash = sha256(output)
    manifest = {
        "schema_version": 1,
        "title": title,
        "research_plan_version": "2.0",
        "generator_version": GENERATOR_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_path": relpath(source),
        "source_sha256": source_hash,
        "source_block_count": block_count,
        "git_commit": commit,
        "git_dirty": dirty,
        "output_path": relpath(output),
        "output_sha256": output_hash,
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
